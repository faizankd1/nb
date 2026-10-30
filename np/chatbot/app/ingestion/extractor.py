from pathlib import Path
import json, csv, shutil, xml.etree.ElementTree as ET
SUPPORTED={".pdf",".docx",".txt",".md",".log",".html",".htm",".rtf",".odt",".jpg",".jpeg",".png",".tif",".tiff",".bmp",".webp",
           ".xlsx",".xlsm",".ods",".csv",".pptx",".json",".xml",".yaml",".yml",".eml"}

class NoReadableTextError(ValueError):
    """The file was processed successfully but contains no indexable text."""


def _get_tesseract():
    from app.core.config import settings
    if not settings.ocr_enabled:
        return None
    cmd = settings.tesseract_cmd
    if cmd and Path(cmd).is_file():
        import pytesseract
        pytesseract.pytesseract.tesseract_cmd = cmd
        return pytesseract
    if shutil.which("tesseract"):
        import pytesseract
        return pytesseract
    return None

def _row_text(row):
    return " | ".join("" if value is None else str(value) for value in row)


def _csv_blocks(path, rows_per_block):
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        header = next(reader, None)
        if header is None:
            return
        batch, start = [], 2
        for row_number, row in enumerate(reader, 2):
            batch.append(f"Row {row_number}: {_row_text(row)}")
            if len(batch) >= rows_per_block:
                yield (f"CSV rows {start}-{row_number}", "Columns: " + _row_text(header) + "\n" + "\n".join(batch))
                batch, start = [], row_number + 1
        if batch:
            yield (f"CSV rows {start}-{start + len(batch) - 1}", "Columns: " + _row_text(header) + "\n" + "\n".join(batch))


def _excel_blocks(path, ext, rows_per_block):
    # read_only mode avoids loading a whole workbook into RAM.  ODS remains on
    # pandas because openpyxl does not support it.
    if ext == ".ods":
        import pandas as pd
        workbook = pd.ExcelFile(path, engine="odf")
        for sheet in workbook.sheet_names:
            frame = pd.read_excel(path, sheet_name=sheet, dtype=str, engine="odf").fillna("")
            header = list(map(str, frame.columns))
            for start in range(0, len(frame), rows_per_block):
                part = frame.iloc[start:start + rows_per_block]
                lines = [f"Row {idx + 2}: " + " | ".join(f"{column}={row[column]}" for column in frame.columns) for idx, row in part.iterrows()]
                if lines:
                    yield (f"sheet {sheet} rows {start + 2}-{start + len(lines) + 1}", "Sheet: " + sheet + "\nColumns: " + " | ".join(header) + "\n" + "\n".join(lines))
        return
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                continue
            batch, start = [], 2
            for row_number, row in enumerate(rows, 2):
                batch.append(f"Row {row_number}: {_row_text(row)}")
                if len(batch) >= rows_per_block:
                    yield (f"sheet {worksheet.title} rows {start}-{row_number}", "Sheet: " + worksheet.title + "\nColumns: " + _row_text(header) + "\n" + "\n".join(batch))
                    batch, start = [], row_number + 1
            if batch:
                yield (f"sheet {worksheet.title} rows {start}-{start + len(batch) - 1}", "Sheet: " + worksheet.title + "\nColumns: " + _row_text(header) + "\n" + "\n".join(batch))
    finally:
        workbook.close()


def iter_extract(path):
    """Yield source-labelled text blocks without loading tabular files as one string."""
    from app.core.config import settings
    ext=Path(path).suffix.lower(); yielded = False
    if ext not in SUPPORTED: raise ValueError(f"Unsupported extension: {ext}")
    if ext==".pdf":
        import pymupdf
        with pymupdf.open(path) as pdf:
            if pdf.page_count > settings.max_pdf_pages:
                raise ValueError(f"PDF has {pdf.page_count} pages, above the configured MAX_PDF_PAGES limit of {settings.max_pdf_pages}.")
            for i,p in enumerate(pdf,1):
                text=p.get_text().strip()
                if not text:
                    pytess = _get_tesseract()
                    if pytess:
                        try:
                            from PIL import Image
                            from io import BytesIO
                            text=pytess.image_to_string(Image.open(BytesIO(p.get_pixmap().tobytes("png")))).strip()
                        except Exception:
                            pass
                if text:
                    yielded = True; yield (f"page {i}",text)
    elif ext==".docx":
        from docx import Document
        d=Document(path); text="\n".join(p.text for p in d.paragraphs if p.text.strip())
        for ti,t in enumerate(d.tables,1):
            text += "\n" + "\n".join(f"[Table {ti}] "+" | ".join(c.text for c in row.cells) for row in t.rows)
        if text.strip(): yielded = True; yield ("document",text)
    elif ext in {".jpg",".jpeg",".png",".tif",".tiff",".bmp",".webp"}:
        pytess = _get_tesseract()
        if not pytess:
            raise RuntimeError(
                f"Tesseract OCR is not installed or configured. Cannot extract text from '{Path(path).name}'. "
                f"Install Tesseract OCR or set TESSERACT_CMD in .env to the tesseract.exe path."
            )
        from PIL import Image, ImageOps
        with Image.open(path) as image:
            if image.width * image.height > settings.max_image_pixels:
                raise ValueError(f"Image has {image.width * image.height:,} pixels, above the configured MAX_IMAGE_PIXELS limit of {settings.max_image_pixels:,}.")
            text = pytess.image_to_string(ImageOps.grayscale(image)).strip()
        if text: yielded = True; yield ("image", text)
    elif ext in {".txt", ".md", ".log", ".html", ".htm", ".yaml", ".yml", ".rtf"}:
        text=Path(path).read_text(encoding="utf-8",errors="replace")
        if ext==".rtf":
            from striprtf.striprtf import rtf_to_text
            text=rtf_to_text(text)
        if text.strip(): yielded = True; yield ("document",text)
    elif ext in {".xlsx",".xlsm",".ods"}:
        for block in _excel_blocks(path, ext, settings.extraction_block_rows):
            yielded = True; yield block
    elif ext==".csv":
        for block in _csv_blocks(path, settings.extraction_block_rows):
            yielded = True; yield block
    elif ext==".pptx":
        from pptx import Presentation
        for i,s in enumerate(Presentation(path).slides,1):
            text="\n".join(sh.text for sh in s.shapes if hasattr(sh,"text") and sh.text.strip())
            if text: yielded = True; yield (f"slide {i}",text)
    elif ext==".json":
        text = json.dumps(json.loads(Path(path).read_text(encoding="utf-8")),ensure_ascii=False,indent=2)
        if text.strip(): yielded = True; yield ("JSON", text)
    elif ext==".xml":
        text = ET.tostring(ET.parse(path).getroot(),encoding="unicode")
        if text.strip(): yielded = True; yield ("XML", text)
    elif ext==".eml":
        from email import policy
        from email.parser import BytesParser
        m=BytesParser(policy=policy.default).parsebytes(Path(path).read_bytes()); body=m.get_body(preferencelist=("plain",))
        text = f"Subject: {m.get('subject','')}\nFrom: {m.get('from','')}\n"+(body.get_content() if body else "")
        if text.strip(): yielded = True; yield ("email", text)
    elif ext==".odt":
        import zipfile
        with zipfile.ZipFile(path) as z: xml=z.read("content.xml")
        tree=ET.fromstring(xml); ns={"t":"urn:oasis:names:tc:opendocument:xmlns:text:1.0"}
        text = "\n".join("".join(p.itertext()) for p in tree.findall(".//t:p",ns))
        if text.strip(): yielded = True; yield ("ODT", text)
    if not yielded:
        raise NoReadableTextError("No readable text was found in the uploaded file. It may be empty, image-only with OCR disabled, or OCR may have failed.")


def extract(path):
    """Compatibility wrapper used by callers that need all blocks at once."""
    return list(iter_extract(path))
